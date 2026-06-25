#!/usr/bin/env python3
"""Convert project agreement Markdown files to .docx (python-docx)."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt


def set_default_font(doc: Document, name: str = "宋体", size_pt: float = 11) -> None:
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(size_pt)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_runs_with_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and "|" in s[1:-1]


def is_table_sep(line: str) -> bool:
    s = line.strip().strip("|")
    if not s:
        return False
    cells = [c.strip() for c in s.split("|")]
    return all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def parse_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell_text = row[ci] if ci < len(row) else ""
            table.rows[ri].cells[ci].text = cell_text
    doc.add_paragraph()


def convert_md_to_docx(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_default_font(doc)

    table_buf: list[list[str]] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if is_table_row(line):
            if is_table_sep(line):
                i += 1
                continue
            table_buf.append(parse_table_row(line))
            i += 1
            continue
        if table_buf:
            flush_table(doc, table_buf)
            table_buf = []

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("—" * 24)
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_heading(stripped[2:].strip(), level=0)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            add_runs_with_bold(p, stripped[2:].strip())
            p.paragraph_format.left_indent = Pt(18)
            i += 1
            continue

        m = re.match(r"^(-|\d+\.)\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet" if m.group(1) == "-" else "List Number")
            add_runs_with_bold(p, m.group(2))
            i += 1
            continue

        if stripped.startswith("|") and "---" in stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs_with_bold(p, stripped)
        i += 1

    if table_buf:
        flush_table(doc, table_buf)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"Wrote {docx_path}")


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print("Usage: md-to-docx.py <file.md> [more.md ...]")
        return 1
    for arg in argv[1:]:
        md_path = Path(arg)
        docx_path = md_path.with_suffix(".docx")
        convert_md_to_docx(md_path, docx_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
